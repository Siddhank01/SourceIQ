import uuid
from pathlib import Path
from typing import List

from openpyxl import load_workbook
from langchain_core.documents import Document
from langchain_community.document_loaders import PyPDFLoader, TextLoader, WebBaseLoader


class DocumentLoader:
    """Load local files and web URLs into LangChain documents."""

    @staticmethod
    def _with_source_metadata(document: Document, source_name: str, source_type: str) -> Document:
        source_id = document.metadata.get("source_id") if isinstance(document.metadata, dict) else None
        if not source_id:
            source_id = str(uuid.uuid4())
        metadata = dict(document.metadata or {})
        metadata.setdefault("source", source_name)
        metadata.setdefault("source_id", source_id)
        metadata.setdefault("source_type", source_type)
        metadata.setdefault("chunk_id", f"{source_id}-{len(metadata.get('chunk_id', ''))}")
        return Document(page_content=document.page_content, metadata=metadata)

    def load_file(self, file_path: str) -> List[Document]:
        source_name = Path(file_path).name
        if file_path.lower().endswith(".pdf"):
            loader = PyPDFLoader(file_path)
            return [self._with_source_metadata(doc, source_name, "pdf") for doc in loader.load()]
        if file_path.lower().endswith(".xlsx"):
            return self._load_workbook(file_path)
        loader = TextLoader(file_path)
        return [self._with_source_metadata(doc, source_name, "text") for doc in loader.load()]

    def _load_workbook(self, file_path: str) -> List[Document]:
        workbook = load_workbook(file_path, read_only=True, data_only=True)
        source = Path(file_path).name
        documents: List[Document] = []
        try:
            for worksheet in workbook.worksheets:
                rows = [
                    "\t".join("" if value is None else str(value) for value in row)
                    for row in worksheet.iter_rows(values_only=True)
                ]
                content = "\n".join(row for row in rows if row.strip())
                if content:
                    document = Document(page_content=content, metadata={"source": source, "sheet": worksheet.title})
                    documents.append(self._with_source_metadata(document, source, "xlsx"))
        finally:
            workbook.close()
        return documents

    def load_url(self, url: str) -> List[Document]:
        loader = WebBaseLoader(url)
        docs = loader.load()
        return [self._with_source_metadata(doc, url, "web") for doc in docs]

    def load_text(self, text: str) -> List[Document]:
        document = Document(page_content=text, metadata={"source": "text_input"})
        return [self._with_source_metadata(document, "text_input", "text")]
